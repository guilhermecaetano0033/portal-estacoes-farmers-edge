"""Gera um XLSX horário completo por cliente para o GitHub Pages.
O arquivo é gerado no runner e fica somente no site publicado, nunca no Git.
"""
import gzip
import json
import logging
from pathlib import Path
from openpyxl import Workbook

from offline_utils import is_offline

ROOT=Path(__file__).resolve().parent
SITE=ROOT/'site'
HEAD=['Estação','ID estação','Data/hora local','Temperatura média (°C)','Temperatura mínima (°C)','Temperatura máxima (°C)','Umidade média (%)','Umidade mínima (%)','Umidade máxima (%)','Chuva (mm)','Vento médio (km/h)','Vento máximo (km/h)','Rajada (km/h)','Direção vento (°)','Radiação solar (MJ/m²)']
KEYS=['temp_media_c','temp_min_c','temp_max_c','umidade_media_pct','umidade_min_pct','umidade_max_pct','chuva_mm','vento_medio_kmh','vento_max_kmh','rajada_kmh','direcao_vento_graus','radiacao_solar_mjm2']
LIMIT=1_000_000

def records(root):
    for path in sorted(root.glob('????-??.json*')):
        try:
            if path.suffix=='.gz':
                with gzip.open(path,'rt',encoding='utf-8') as h:
                    for row in json.load(h)['rows']:
                        if not is_offline(row):
                            yield row
            elif not path.with_suffix(path.suffix+'.gz').exists():
                with path.open(encoding='utf-8') as h:
                    for row in json.load(h)['rows']:
                        if not is_offline(row):
                            yield row
        except (OSError,ValueError,KeyError) as e:
            raise RuntimeError(f'Falha lendo {path} ({type(e).__name__})') from e

def workbook(client, target):
    """Cria um único Excel horário consolidado; escrita incremental e pouco uso de memória."""
    wb=Workbook(write_only=True)
    num=1
    ws=wb.create_sheet('Horarios 1')
    ws.append(HEAD)
    total=0
    for st in client['estacoes']:
        for r in records(SITE/'data'/client['slug']/str(st['id'])):
            if total and total % LIMIT == 0:
                num+=1
                ws=wb.create_sheet('Horarios '+str(num))
                ws.append(HEAD)
            timestamp=str(r.get('data_hora_local','')).replace('T',' ')[:16]
            if not timestamp:continue
            ws.append([st['nome'],st['id'],timestamp]+[r.get(k) for k in KEYS])
            total+=1
    coords=wb.create_sheet('Coordenadas')
    coords.append(['Estação','ID estação','FEID','Latitude','Longitude','Cliente'])
    for st in client['estacoes']:
        coords.append([st['nome'],st['id'],st.get('feid'),st.get('latitude'),st.get('longitude'),client['nome']])
    target.parent.mkdir(parents=True,exist_ok=True)
    wb.save(target)
    return total

def main():
    from tempfile import TemporaryDirectory
    cfg=json.loads((ROOT/'clientes.json').read_text(encoding='utf-8'))
    folder=SITE/'downloads';folder.mkdir(exist_ok=True)
    for client in cfg['clientes']:
        name=client['slug']+'_historico_horario.xlsx'
        final=folder/name
        with TemporaryDirectory(dir=folder) as tmp:
            temp=Path(tmp)/name
            n=workbook(client,temp)
            temp.replace(final)
        logging.info('%s: %s horários válidos em %s (%.2f MB)',client['slug'],n,final,final.stat().st_size/1048576)

if __name__=='__main__':
    logging.basicConfig(level=logging.INFO)
    main()
