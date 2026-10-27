import gzip
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
import gerar_downloads

ROOT=Path(__file__).resolve().parents[1]
EXPECTED={
    'tessaro':{39266:'FSE09',37191:'FSE03',80986:'FSE04',80977:'FSE05',80985:'FSE07',65610:'FSE11',80984:'FSE12',59504:'FSE13'},
    'fenix':{81109:'FEX02',81110:'FRO01',81111:'FUA07',81112:'FEX01',81113:'FJG01',81254:'FRO02',81255:'FJG02',81256:'FUA08',81257:'FUA01',81263:'FEX03'},
}
class Version9Tests(unittest.TestCase):
    def test_all_18_stations_renamed_by_stable_id(self):
        data=json.loads((ROOT/'clientes.json').read_text(encoding='utf8'))
        self.assertEqual(set(c['slug'] for c in data['clientes']),set(EXPECTED))
        for client in data['clientes']:
            self.assertEqual({s['id']:s['nome'] for s in client['estacoes']},EXPECTED[client['slug']])
            self.assertTrue(all('latitude' in s and 'longitude' in s and 'feid' in s for s in client['estacoes']))

    def test_pages_only_have_two_excel_buttons(self):
        for slug in EXPECTED:
            text=(ROOT/'site'/f'{slug}.html').read_text(encoding='utf8')
            self.assertIn('id="periodExcel"',text)
            self.assertIn('id="fullExcel"',text)
            for old in ['id="rename"','id="renameInfo"','id="csv"','id="zip"','id="allzip"','id="fullzip"']:
                self.assertNotIn(old,text)

    def test_full_history_single_hourly_xlsx_and_coordinates_offline_filtered(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp=Path(tmp)
            month=tmp/'data'/'tessaro'/'39266'/'2026-10.json.gz'
            month.parent.mkdir(parents=True)
            rows=[
                {'data_hora_local':'2026-10-01T00:00:00-04:00','temp_media_c':None,'umidade_media_pct':None,'chuva_mm':2.0},
                {'data_hora_local':'2026-10-01T01:00:00-04:00','temp_media_c':24.1,'umidade_media_pct':70.0,'chuva_mm':0},
                {'data_hora_local':'2026-10-01T02:00:00-04:00','temp_media_c':24.5,'umidade_media_pct':None,'chuva_mm':2.3},
            ]
            with gzip.open(month,'wt',encoding='utf8') as f:
                json.dump({'rows':rows},f)
            data=json.loads((ROOT/'clientes.json').read_text(encoding='utf8'))
            client=next(c for c in data['clientes'] if c['slug']=='tessaro')
            target=tmp/'export.xlsx'
            with patch.object(gerar_downloads,'SITE',tmp):
                n=gerar_downloads.workbook(client,target)
            self.assertEqual(n,2)
            wb=load_workbook(target,read_only=True,data_only=True)
            self.assertEqual(wb.sheetnames,['Horarios 1','Coordenadas'])
            vals=list(wb['Horarios 1'].values)
            self.assertEqual(len(vals),3)
            self.assertEqual(vals[1][0],'FSE09')
            self.assertEqual(vals[1][2],'2026-10-01 01:00')
            self.assertEqual(vals[2][2],'2026-10-01 02:00')
            coords=list(wb['Coordenadas'].values)
            self.assertEqual(len(coords),9)
            self.assertEqual(coords[1][0],'FSE09')
            self.assertEqual(coords[1][3],-12.4048)
            self.assertNotIn('Diarios 1',wb.sheetnames)
            wb.close()

if __name__=='__main__':unittest.main()
