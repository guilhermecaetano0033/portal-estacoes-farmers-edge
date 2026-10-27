import gzip
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import atualizar
import buscar_talhoes
import gerar_downloads
from offline_utils import is_offline
from openpyxl import load_workbook


class OfflineTests(unittest.TestCase):
    def test_missing_temperature_and_humidity_means_offline_even_with_rain(self):
        row = atualizar.normalize({'local_time': '2026-10-20T11:00:00-04:00', 'total_precip_mm': 12}, 81263)
        self.assertTrue(row['offline'])
        self.assertTrue(is_offline(row))

    def test_temperature_alone_keeps_record(self):
        row = atualizar.normalize({'local_time': '2026-10-20T12:00:00-04:00', 'avg_temp_c': 24}, 81263)
        self.assertFalse(row['offline'])
        self.assertFalse(is_offline(row))

    def test_zero_is_an_observation_not_missing(self):
        self.assertFalse(is_offline({'temp_media_c': 0, 'umidade_media_pct': None}))
        self.assertTrue(is_offline({'temp_media_c': None, 'umidade_media_pct': None}))

    def test_export_filters_legacy_rows_without_status_column(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            folder = tmp / 'data' / 'x' / '100'
            folder.mkdir(parents=True)
            path = folder / '2026-10.json.gz'
            rows = [
                {'data_hora_local': '2026-10-01T00:00:00-04:00', 'temp_media_c': None, 'umidade_media_pct': None, 'chuva_mm': 20},
                {'data_hora_local': '2026-10-01T01:00:00-04:00', 'temp_media_c': 25, 'umidade_media_pct': 60, 'chuva_mm': 2},
            ]
            with gzip.open(path, 'wt', encoding='utf8') as h:
                json.dump({'rows': rows}, h)
            self.assertEqual(len(list(gerar_downloads.records(folder))), 1)
            client = {'slug': 'x', 'nome': 'X', 'estacoes': [{'id': 100, 'nome': 'E100', 'latitude': -12, 'longitude': -56}]}
            with patch.object(gerar_downloads, 'SITE', tmp):
                hour = tmp / 'hour.xlsx'
                self.assertEqual(gerar_downloads.workbook(client, hour), 1)
            book = load_workbook(hour, read_only=True)
            self.assertEqual(len(list(book['Horarios 1'].rows)), 2)
            self.assertIn('Coordenadas', book.sheetnames)
            self.assertFalse(any('Diarios' in n for n in book.sheetnames))
            book.close()


class FieldBordersTests(unittest.TestCase):
    def test_only_fields_owned_by_grower(self):
        assets = [
            {'id': 92088, 'category': 'Grower'},
            {'id': 8, 'category': 'Farm', 'parent': {'id': 92088}},
            {'id': 11, 'category': 'Field', 'parent': 8},
            {'id': 12, 'category': 'Field', 'parent': 8},
            {'id': 99, 'category': 'Field', 'parent': 12345},
        ]
        self.assertEqual({f['id'] for f in buscar_talhoes.fields_for_grower(assets, 92088)}, {11, 12})

    def test_feature_collection_and_multipolygon_are_preserved(self):
        poly = {'type': 'Polygon', 'coordinates': [[[-56.0, -12.0], [-55.99, -12.0], [-55.99, -12.01], [-56.0, -12.0]]]}
        multipoly = {'type': 'MultiPolygon', 'coordinates': [poly['coordinates']]}
        result = {'results': [{'shapeData': json.dumps({'type': 'FeatureCollection', 'features': [
            {'type': 'Feature', 'geometry': poly}, {'type': 'Feature', 'geometry': multipoly}
        ]})}]}
        self.assertEqual(len(list(buscar_talhoes.geometries(result))), 2)


if __name__ == '__main__':
    unittest.main()
